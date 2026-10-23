import re
from typing import Any
import pypdf
import docx

# Normalized technology skill catalog with common variations
SKILL_TAXONOMY = {
    "Python": ["python", "python3", "py"],
    "JavaScript": ["javascript", "js", "ecmascript"],
    "TypeScript": ["typescript", "ts"],
    "React": ["react", "react.js", "reactjs"],
    "Next.js": ["next.js", "nextjs", "next"],
    "Node.js": ["node.js", "nodejs", "node"],
    "FastAPI": ["fastapi", "fast-api"],
    "Express.js": ["express", "express.js", "expressjs"],
    "Django": ["django"],
    "Flask": ["flask"],
    "HTML5": ["html", "html5"],
    "CSS3": ["css", "css3", "tailwind", "tailwindcss", "sass", "bootstrap"],
    "Tailwind CSS": ["tailwind", "tailwindcss"],
    "SQL": ["sql", "mysql", "mariadb"],
    "PostgreSQL": ["postgresql", "postgres", "psql"],
    "MongoDB": ["mongodb", "mongo"],
    "Redis": ["redis"],
    "Git": ["git", "github", "gitlab"],
    "Docker": ["docker", "containerization", "containers"],
    "Kubernetes": ["kubernetes", "k8s"],
    "AWS": ["aws", "amazon web services", "ec2", "s3", "lambda"],
    "Azure": ["azure", "microsoft azure"],
    "GCP": ["gcp", "google cloud platform", "google cloud"],
    "Linux": ["linux", "ubuntu", "debian", "bash", "shell scripting"],
    "REST APIs": ["rest api", "restful api", "rest apis", "restful apis", "rest"],
    "GraphQL": ["graphql"],
    "CI/CD": ["ci/cd", "continuous integration", "github actions", "jenkins"],
    "Machine Learning": ["machine learning", "ml", "scikit-learn", "sklearn"],
    "Deep Learning": ["deep learning", "neural networks", "pytorch", "tensorflow", "keras"],
    "Data Analysis": ["data analysis", "pandas", "numpy", "data visualization", "matplotlib", "seaborn"],
    "Cybersecurity Fundamentals": ["cybersecurity", "infosec", "network security", "penetration testing", "owasp"],
    "Testing & QA": ["testing", "unit testing", "pytest", "jest", "cypress", "playwright"],
    "Data Structures & Algorithms": ["data structures", "algorithms", "dsa", "leetcode"],
    "System Design": ["system design", "software architecture", "microservices"],
    "Agile / Scrum": ["agile", "scrum", "kanban", "sprint"],
    "UI/UX Design": ["ui/ux", "figma", "wireframing", "prototyping", "user research"],
}

def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Extract raw text from PDF, DOCX, or TXT."""
    text = ""
    try:
        if file_type.endswith("pdf"):
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        elif file_type.endswith("docx") or file_type.endswith("doc"):
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                    text += "\n"
        else:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
    except Exception as e:
        text = f"Error extracting document text: {str(e)}"
    return text.strip()

def extract_sections(text: str) -> dict[str, list[str]]:
    """Heuristic extraction of education, experience, projects, and skills."""
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    
    sections = {
        "education": [],
        "experience": [],
        "projects": [],
        "skills_raw": [],
        "general": []
    }
    
    current_section = "general"
    
    section_patterns = {
        "education": re.compile(r"^(education|academic background|qualifications|university|degrees)", re.I),
        "experience": re.compile(r"^(experience|work experience|employment history|internships|work history)", re.I),
        "projects": re.compile(r"^(projects|academic projects|personal projects|portfolio projects)", re.I),
        "skills_raw": re.compile(r"^(skills|technical skills|technologies|competencies|tools & technologies)", re.I)
    }
    
    for line in lines:
        matched = False
        for sec, pattern in section_patterns.items():
            if pattern.match(line) and len(line) < 35:
                current_section = sec
                matched = True
                break
        if not matched:
            sections[current_section].append(line)
            
    return sections

def extract_skills_from_text(text: str) -> list[str]:
    """Finds canonical skills mentioned in text with word boundary matching."""
    lower_text = " " + text.lower() + " "
    detected = set()
    
    for canonical, aliases in SKILL_TAXONOMY.items():
        for alias in aliases:
            # Word boundary regex to avoid substring false positives (e.g. "c" in "cat")
            pattern = rf"(?:\b|[^a-zA-Z0-9]){re.escape(alias)}(?:\b|[^a-zA-Z0-9])"
            if re.search(pattern, lower_text):
                detected.add(canonical)
                break
                
    return sorted(list(detected))

def generate_cv_feedback(extracted_skills: list[str], sections: dict[str, list[str]], target_career_slug: str, career_requirements: list[Any]) -> tuple[list[str], list[str], list[str]]:
    """Generates structured CV strengths, missing critical skills, and improvements."""
    strengths = []
    missing_critical = []
    feedback = []
    
    # 1. Evaluate section completeness
    if not sections["education"]:
        feedback.append("No distinct Education section found. Ensure your university degree, institution, and graduation year are clearly listed.")
    else:
        strengths.append("Education section is clearly identifiable with academic credentials.")
        
    if not sections["projects"]:
        feedback.append("No dedicated Projects section identified. For technology roles, practical project demonstrations with GitHub links are essential.")
    else:
        strengths.append(f"Contains project demonstrations showcasing applied competencies.")

    if not sections["experience"]:
        feedback.append("No prior formal Work Experience or Internship section detected. Highlight leadership, lab demonstrations, or capstone projects as experience.")
    else:
        strengths.append("Demonstrates documented practical internship or workplace experience.")

    # 2. Check for metric-driven phrasing (bullet points with numbers/impact)
    all_text = " ".join([" ".join(v) for v in sections.values()])
    has_metrics = bool(re.search(r"\b(\d+%\b|\$\d+|\d+\+?\s*(users|clients|students|queries|requests|endpoints|stars))\b", all_text, re.I))
    if has_metrics:
        strengths.append("Uses quantifiable metrics and numerical impact to describe achievements.")
    else:
        feedback.append("Enhance bullet points with measurable impact (e.g., 'Optimized query latency by 35%' or 'Served 500+ active users').")

    # 3. Match against target career requirements
    detected_set = set(extracted_skills)
    for req in career_requirements:
        req_name = req.skill.name if hasattr(req, "skill") and req.skill else req.get("skill_name", "")
        importance = getattr(req, "importance", req.get("importance", "required") if isinstance(req, dict) else "required")
        if req_name and req_name not in detected_set:
            if importance == "required":
                missing_critical.append(req_name)

    if missing_critical:
        feedback.append(f"Target role requirements not found in CV: {', '.join(missing_critical[:5])}. Consider including course projects demonstrating these.")
    else:
        strengths.append("Contains coverage of foundational requirements for the targeted role.")

    return strengths, missing_critical, feedback
